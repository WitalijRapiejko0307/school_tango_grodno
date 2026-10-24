"""Excel group attendance sheet generation."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    Attendance,
    Group,
    GroupMembership,
    Payment,
    Person,
    SchoolSession,
    Subscription,
)
from app.services.school_time import school_tz


def _session_local_date(session_row: SchoolSession) -> date:
    if session_row.session_date is not None:
        return session_row.session_date
    starts = session_row.starts_at
    if starts.tzinfo is None:
        return starts.replace(tzinfo=school_tz()).date()
    return starts.astimezone(school_tz()).date()


def _couple_display_name(name1: str, name2: str) -> str:
    parts1 = name1.split()
    parts2 = name2.split()
    if parts1 and parts2 and parts1[0] == parts2[0]:
        rest1 = " ".join(parts1[1:])
        rest2 = " ".join(parts2[1:])
        if rest1 and rest2:
            return f"{parts1[0]} {rest1} и {rest2}"
        if rest1:
            return f"{parts1[0]} {rest1} и {name2}"
        if rest2:
            return f"{name1} и {rest2}"
        return f"{name1} и {name2}"
    return f"{name1} и {name2}"


def _row_display_name(person_ids: tuple[str, ...], persons: dict[str, Person]) -> str:
    if len(person_ids) == 1:
        return persons[person_ids[0]].full_name
    a, b = person_ids[0], person_ids[1]
    n1, n2 = persons[a].full_name, persons[b].full_name
    if n1.casefold() <= n2.casefold():
        return _couple_display_name(n1, n2)
    return _couple_display_name(n2, n1)


def _subscription_covers(sub: Subscription, session_day: date) -> bool:
    if sub.starts_on is not None and session_day < sub.starts_on:
        return False
    if sub.valid_until is not None and session_day > sub.valid_until:
        return False
    return True


def _pick_covering_subscription(
    subscriptions: list[Subscription], session_day: date
) -> Subscription | None:
    covering = [s for s in subscriptions if _subscription_covers(s, session_day)]
    if not covering:
        return None
    return max(covering, key=lambda s: (s.starts_on is not None, s.starts_on or date.min))


_WEEKDAYS = ("пн", "вт", "ср", "чт", "пт", "сб", "вс")

_VERTICAL = Alignment(textRotation=90, horizontal="center", vertical="center", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="center")
_NAME = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _header_date(session_day: date) -> str:
    return f"{session_day.strftime('%d.%m')} {_WEEKDAYS[session_day.weekday()]}"


def _fit_sheet(ws, name_column_width: float, date_column_count: int) -> None:
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = max(name_column_width, 22)
    for offset in range(date_column_count):
        letter = get_column_letter(3 + offset)
        ws.column_dimensions[letter].width = 5
    ws.row_dimensions[1].height = 58
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 18
    ws.sheet_view.showGridLines = True
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _payment_total_text(payments: list[Payment]) -> str | None:
    if not payments:
        return None
    total = sum((p.amount for p in payments), Decimal("0"))
    return f"{total:.2f} руб."


async def build_group_attendance_workbook(
    session: AsyncSession,
    group_id: str,
    date_from: date,
    date_to: date,
) -> bytes | None:
    group = await session.get(Group, group_id)
    if group is None:
        return None

    sessions_result = await session.execute(
        select(SchoolSession)
        .where(
            SchoolSession.group_id == group_id,
            SchoolSession.status != "cancelled",
        )
        .order_by(SchoolSession.starts_at)
    )
    all_sessions = list(sessions_result.scalars().all())
    column_sessions = [
        s
        for s in all_sessions
        if date_from <= _session_local_date(s) <= date_to
    ]
    if not column_sessions:
        return None

    memberships_result = await session.execute(
        select(GroupMembership).where(
            GroupMembership.group_id == group_id,
            GroupMembership.started_on <= date_to,
            or_(
                GroupMembership.ended_on.is_(None),
                GroupMembership.ended_on >= date_from,
            ),
        )
    )
    memberships = list(memberships_result.scalars().all())
    member_person_ids = {m.person_id for m in memberships}
    if not member_person_ids:
        wb = Workbook()
        ws = wb.active
        ws.title = "Посещаемость"
        ws["A1"] = "№ п.п."
        ws["A1"].alignment = _VERTICAL
        ws["A1"].font = Font(bold=True)
        ws["B1"] = f"{group.name}\nФ. И."
        ws["B1"].alignment = _NAME
        ws["B1"].font = Font(bold=True)
        for col_idx, school_session in enumerate(column_sessions, start=3):
            cell = ws.cell(
                row=1,
                column=col_idx,
                value=_header_date(_session_local_date(school_session)),
            )
            cell.alignment = _VERTICAL
            cell.font = Font(bold=True)
        _fit_sheet(ws, 22, len(column_sessions))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    persons_result = await session.execute(
        select(Person).where(Person.id.in_(member_person_ids))
    )
    persons = {p.id: p for p in persons_result.scalars().all()}

    subs_result = await session.execute(
        select(Subscription).where(
            or_(
                Subscription.person_id.in_(member_person_ids),
                Subscription.partner_id.in_(member_person_ids),
            ),
        )
    )
    all_subscriptions = list(subs_result.scalars().all())
    couple_pairs: set[frozenset[str]] = set()
    for sub in all_subscriptions:
        if sub.partner_id is None:
            continue
        pair = frozenset({sub.person_id, sub.partner_id})
        if pair <= member_person_ids:
            couple_pairs.add(pair)

    remaining = set(member_person_ids)
    row_units: list[tuple[str, ...]] = []
    for pair in couple_pairs:
        ids = tuple(pair)
        if ids[0] in remaining and ids[1] in remaining:
            row_units.append(ids)
            remaining -= set(ids)
    for pid in remaining:
        row_units.append((pid,))
    row_units.sort(key=lambda u: _row_display_name(u, persons).casefold())

    involved_person_ids = set(member_person_ids)

    attendances_result = await session.execute(
        select(Attendance, SchoolSession)
        .join(SchoolSession, Attendance.session_id == SchoolSession.id)
        .where(Attendance.person_id.in_(involved_person_ids))
    )
    attendance_rows = list(attendances_result.all())
    sessions_by_id: dict[str, SchoolSession] = {s.id: s for s in column_sessions}
    for _att, school_session in attendance_rows:
        sessions_by_id[school_session.id] = school_session

    row_attended_sessions: dict[tuple[str, ...], set[str]] = {}
    for unit in row_units:
        unit_set = set(unit)
        attended: set[str] = set()
        for att, school_session in attendance_rows:
            if att.person_id not in unit_set:
                continue
            attended.add(school_session.id)
        row_attended_sessions[unit] = attended

    payments_by_sub: dict[str, list[Payment]] = {}
    if all_subscriptions:
        sub_ids = {s.id for s in all_subscriptions}
        payments_result = await session.execute(
            select(Payment).where(Payment.subscription_id.in_(sub_ids))
        )
        for payment in payments_result.scalars().all():
            payments_by_sub.setdefault(payment.subscription_id, []).append(payment)

    subs_by_person: dict[str, list[Subscription]] = {}
    for sub in all_subscriptions:
        for pid in (sub.person_id, sub.partner_id):
            if pid is None:
                continue
            if pid in involved_person_ids:
                subs_by_person.setdefault(pid, []).append(sub)

    def subscriptions_for_unit(unit: tuple[str, ...]) -> list[Subscription]:
        seen: set[str] = set()
        out: list[Subscription] = []
        for pid in unit:
            for sub in subs_by_person.get(pid, []):
                if sub.id not in seen:
                    seen.add(sub.id)
                    out.append(sub)
        return out

    def ordinal_for_session(
        unit: tuple[str, ...],
        school_session: SchoolSession,
        unit_subs: list[Subscription],
    ) -> tuple[int | None, Subscription | None]:
        session_day = _session_local_date(school_session)
        if school_session.id not in row_attended_sessions.get(unit, set()):
            return None, None
        sub = _pick_covering_subscription(unit_subs, session_day)
        if sub is None:
            return None, None

        unit_set = set(unit)
        in_window: list[str] = []
        for att, sess in attendance_rows:
            if att.person_id not in unit_set:
                continue
            day = _session_local_date(sess)
            if not _subscription_covers(sub, day):
                continue
            in_window.append(sess.id)

        unique_ids = list(dict.fromkeys(in_window))
        unique_ids.sort(key=lambda sid: sessions_by_id[sid].starts_at)
        try:
            pos = unique_ids.index(school_session.id) + 1
        except ValueError:
            return None, None
        return pos, sub

    wb = Workbook()
    ws = wb.active
    ws.title = "Посещаемость"
    ws["A1"] = "№ п.п."
    ws["A1"].alignment = _VERTICAL
    ws["A1"].font = Font(bold=True)
    ws["B1"] = f"{group.name}\nФ. И."
    ws["B1"].alignment = _NAME
    ws["B1"].font = Font(bold=True)

    for col_idx, school_session in enumerate(column_sessions, start=3):
        cell = ws.cell(
            row=1,
            column=col_idx,
            value=_header_date(_session_local_date(school_session)),
        )
        cell.alignment = _VERTICAL
        cell.font = Font(bold=True)

    longest_name = len(f"{group.name} Ф. И.")
    for row_idx, unit in enumerate(row_units, start=2):
        number = ws.cell(row=row_idx, column=1, value=row_idx - 1)
        number.alignment = _CENTER
        display = _row_display_name(unit, persons)
        longest_name = max(longest_name, len(display))
        name_cell = ws.cell(row=row_idx, column=2, value=display)
        name_cell.alignment = _NAME
        unit_subs = subscriptions_for_unit(unit)
        for col_idx, school_session in enumerate(column_sessions, start=3):
            ordinal, sub = ordinal_for_session(unit, school_session, unit_subs)
            if ordinal is None:
                continue
            cell = ws.cell(row=row_idx, column=col_idx, value=ordinal)
            cell.alignment = _CENTER
            if ordinal == 1 and sub is not None:
                comment_text = _payment_total_text(payments_by_sub.get(sub.id, []))
                if comment_text is not None:
                    cell.comment = Comment(comment_text, "")

    _fit_sheet(ws, longest_name + 2, len(column_sessions))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
