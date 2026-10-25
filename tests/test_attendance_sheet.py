from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.db import Base
from app.models import Person
from app.services.attendance import mark_attendance
from app.services.attendance_sheet import build_group_attendance_workbook
from app.services.billing import create_product, open_subscription, set_price
from app.services.schedule import assign_group, create_group, create_session

MINSK = ZoneInfo("Europe/Minsk")


@pytest.fixture
async def db_session() -> AsyncSession:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    session_factory = async_sessionmaker(
        engine, class_=AsyncSession, expire_on_commit=False
    )
    async with session_factory() as session:
        yield session
    await engine.dispose()


async def test_visit_numbers_restart_on_next_subscription(db_session: AsyncSession) -> None:
    person = Person(full_name="Иванов Иван", role="client", telegram_user_id=501)
    group = await create_group(db_session, "Начинающие")
    db_session.add(person)
    await db_session.flush()
    await assign_group(db_session, person.id, group.id, date(2026, 9, 1))
    drop_in = await create_product(db_session, "drop_in", "Разовое")
    await set_price(db_session, drop_in.id, Decimal("15.00"), date(2026, 1, 1))
    product = await create_product(
        db_session, "subscription", "8 занятий", lessons_count=8, validity_days=28
    )
    first = await open_subscription(
        db_session,
        person.id,
        product.id,
        started_on=date(2026, 9, 1),
        starts_on=date(2026, 9, 1),
        amount_due=Decimal("100.00"),
        amount_paid=Decimal("100.00"),
    )
    first.valid_until = date(2026, 9, 10)
    second = await open_subscription(
        db_session,
        person.id,
        product.id,
        started_on=date(2026, 9, 11),
        starts_on=date(2026, 9, 11),
        amount_due=Decimal("100.00"),
        amount_paid=Decimal("40.00"),
    )
    second.valid_until = date(2026, 9, 30)
    days = [date(2026, 9, 2), date(2026, 9, 4), date(2026, 9, 12), date(2026, 9, 14)]
    sessions = []
    for day in days:
        starts = datetime(day.year, day.month, day.day, 19, 0, tzinfo=MINSK)
        sessions.append(
            await create_session(
                db_session,
                group.id,
                starts_at=starts,
                ends_at=starts + timedelta(hours=1),
                place="Зал",
                bring_notes=None,
            )
        )
    outside = datetime(2026, 8, 20, 19, 0, tzinfo=MINSK)
    early = await create_session(
        db_session,
        group.id,
        starts_at=outside,
        ends_at=outside + timedelta(hours=1),
        place="Зал",
        bring_notes=None,
    )
    for school_session in [*sessions, early]:
        await mark_attendance(
            db_session, person.id, school_session.id, "admin", person.id
        )
    await db_session.flush()

    payload = await build_group_attendance_workbook(
        db_session, group.id, date(2026, 9, 1), date(2026, 9, 30)
    )
    assert payload is not None
    wb = load_workbook(BytesIO(payload))
    ws = wb.active
    assert [ws.cell(1, col).value for col in range(3, 7)] == [
        "02.09 ср",
        "04.09 пт",
        "12.09 сб",
        "14.09 пн",
    ]
    assert ws.cell(1, 1).alignment.textRotation == 90
    assert ws.cell(1, 3).alignment.textRotation == 90
    assert ws.cell(2, 1).alignment.textRotation == 0
    assert [ws.cell(2, col).value for col in range(3, 7)] == [1, 2, 1, 2]
    assert ws.cell(2, 3).comment.content.startswith("100.00 руб.")
    assert ws.cell(2, 5).comment.content.startswith("40.00 руб.")
    assert ws.cell(2, 4).comment is None
